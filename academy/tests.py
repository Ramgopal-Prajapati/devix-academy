from django.test import TestCase, Client
from django.urls import reverse
from decimal import Decimal
from datetime import date
from accounts.models import User
from academy.models import Department, Course, TrainerProfile, StudentProfile, generate_student_id
from attendance.models import Attendance


class AcademyFunctionalityTests(TestCase):
    def setUp(self):
        self.client = Client()
        self.dept = Department.objects.create(name="Tech", description="Technology courses")
        self.course = Course.objects.create(
            name="Python Full Stack",
            department=self.dept,
            description="Django & React",
            duration="6 Months",
            fees=Decimal("35000.00")
        )

        self.bh_user = User.objects.create_user(
            username='admin_ram',
            email='ram@devix.com',
            password='Password123!',
            role=User.Role.BRANCH_HEAD
        )

        self.trainer_user = User.objects.create_user(
            username='trainer_mark',
            email='mark@devix.com',
            password='Password123!',
            role=User.Role.TRAINER
        )
        self.trainer = TrainerProfile.objects.create(
            user=self.trainer_user,
            trainer_id="TRN-901",
            phone="9898989801",
            email="mark@devix.com",
            department=self.dept,
            monthly_salary=Decimal("60000.00")
        )

        self.student_user = User.objects.create_user(
            username='student_clara',
            email='clara@gmail.com',
            password='Password123!',
            role=User.Role.STUDENT
        )
        self.student = StudentProfile.objects.create(
            user=self.student_user,
            student_id="DVX-2026-0901",
            name="Clara Oswald",
            email="clara@gmail.com",
            phone="9898989899",
            department=self.dept,
            course=self.course,
            trainer=self.trainer,
            start_date=date.today()
        )

    def test_public_home_security_no_sensitive_data(self):
        """Verify public home displays student name and course, but NO email or phone."""
        response = self.client.get(reverse('public_home'))
        self.assertEqual(response.status_code, 200)
        content = response.content.decode('utf-8')
        self.assertIn("Clara Oswald", content)
        self.assertIn("Python Full Stack", content)
        self.assertNotIn("clara@gmail.com", content)
        self.assertNotIn("9898989899", content)

    def test_student_id_generator(self):
        """Verify student ID follows DVX-{year}-{0001} sequence."""
        sid = generate_student_id()
        self.assertTrue(sid.startswith("DVX-"))
        parts = sid.split("-")
        self.assertEqual(len(parts), 3)

    def test_certificate_eligibility_logic(self):
        """Verify attendance >= 75% is eligible and < 75% is ineligible."""
        # Record 4 attendance days: 3 Present, 1 Absent (75%)
        Attendance.objects.create(student=self.student, date=date(2026, 1, 1), status='Present')
        Attendance.objects.create(student=self.student, date=date(2026, 1, 2), status='Present')
        Attendance.objects.create(student=self.student, date=date(2026, 1, 3), status='Present')
        Attendance.objects.create(student=self.student, date=date(2026, 1, 4), status='Absent')

        stats = self.student.get_attendance_stats()
        self.assertEqual(stats['percentage'], 75.0)
        self.assertTrue(stats['is_eligible'])

        # Add 2 more Absent days (3 / 6 = 50.0%) -> Ineligible
        Attendance.objects.create(student=self.student, date=date(2026, 1, 5), status='Absent')
        Attendance.objects.create(student=self.student, date=date(2026, 1, 6), status='Absent')

        stats = self.student.get_attendance_stats()
        self.assertEqual(stats['percentage'], 50.0)
        self.assertFalse(stats['is_eligible'])

    def test_courses_api_filtered_by_department(self):
        """Verify /api/courses/?department=<id> filters courses correctly."""
        dept2 = Department.objects.create(name="Non-Tech")
        Course.objects.create(
            name="Graphic Design",
            department=dept2,
            description="Photoshop & Illustrator",
            duration="3 Months"
        )

        res = self.client.get(f'/api/courses/?department={self.dept.id}')
        self.assertEqual(res.status_code, 200)
        data = res.json()
        results = data if isinstance(data, list) else data.get('results', [])
        self.assertEqual(len(results), 1)
        self.assertEqual(results[0]['name'], "Python Full Stack")

    def test_trainer_leave_notice_isolated_to_assigned_students(self):
        """Verify Trainer leave notice is visible only to assigned students."""
        from academy.models import TrainerNotice

        # Notice from Clara's trainer (Mark)
        notice_mark = TrainerNotice.objects.create(
            trainer=self.trainer,
            department=self.dept,
            notice_type='LEAVE',
            title='Mentor on Leave',
            message='I will be out of office tomorrow.'
        )

        # Another trainer & student
        trainer2_user = User.objects.create_user(username='t2', email='t2@devix.com', password='Password123!', role=User.Role.TRAINER)
        trainer2 = TrainerProfile.objects.create(user=trainer2_user, trainer_id='TRN-902', phone='9898989802', email='t2@devix.com')
        student2_user = User.objects.create_user(username='s2', email='s2@gmail.com', password='Password123!', role=User.Role.STUDENT)
        student2 = StudentProfile.objects.create(user=student2_user, student_id='DVX-2026-0902', name='Other Student', email='s2@gmail.com', phone='9898989803', trainer=trainer2)

        # Clara views notices
        self.client.login(username='student_clara', password='Password123!')
        res_clara = self.client.get(reverse('student:notices'))
        self.assertEqual(res_clara.status_code, 200)
        self.assertContains(res_clara, "Mentor on Leave")
        self.client.logout()

        # Student 2 views notices
        self.client.login(username='s2', password='Password123!')
        res_s2 = self.client.get(reverse('student:notices'))
        self.assertEqual(res_s2.status_code, 200)
        self.assertNotContains(res_s2, "Mentor on Leave")
        self.client.logout()

    def test_student_direct_message_delivery(self):
        """Verify trainer can send individual message and materials to a specific student."""
        from academy.models import StudentDirectMessage

        msg = StudentDirectMessage.objects.create(
            trainer=self.trainer,
            student=self.student,
            subject='Specialized Project Guide',
            message='Read this document carefully.',
            link_url='https://github.com/devix'
        )

        self.client.login(username='student_clara', password='Password123!')
        res = self.client.get(reverse('student:messages'))
        self.assertEqual(res.status_code, 200)
        self.assertContains(res, "Specialized Project Guide")
        self.assertContains(res, "Read this document carefully.")
        self.assertContains(res, "https://github.com/devix")

        msg.refresh_from_db()
        self.assertTrue(msg.is_read)
        self.client.logout()

    def test_course_resource_filtered_strictly_by_student_course(self):
        """Verify students only see resources matching their enrolled course."""
        from academy.models import CourseResource

        course2 = Course.objects.create(name="Data Analytics", department=self.dept, description="DA track", duration="4 Months")

        # Resource for Python Full Stack
        res1 = CourseResource.objects.create(
            trainer=self.trainer,
            department=self.dept,
            course=self.course,
            title='Python Full Stack CheatSheet',
            resource_type='NOTES',
            notes_content='print("Hello World")'
        )

        # Resource for Data Analytics
        res2 = CourseResource.objects.create(
            trainer=self.trainer,
            department=self.dept,
            course=course2,
            title='Data Analytics Excel Formula Book',
            resource_type='LINK',
            link_url='https://excel.com'
        )

        # Clara (enrolled in Python Full Stack) visits resources
        self.client.login(username='student_clara', password='Password123!')
        res_page = self.client.get(reverse('student:resources'))
        self.assertEqual(res_page.status_code, 200)
        self.assertContains(res_page, "Python Full Stack CheatSheet")
        self.assertNotContains(res_page, "Data Analytics Excel Formula Book")
        self.client.logout()

    def test_two_way_student_trainer_1_on_1_chat(self):
        """Verify bidirectional 1-on-1 private messaging between student and trainer."""
        from academy.models import StudentDirectMessage

        # Trainer sends initial message
        msg = StudentDirectMessage.objects.create(
            trainer=self.trainer,
            student=self.student,
            sender_role='TRAINER',
            subject='Initial Guidance',
            message='Please start working on Project 1.'
        )

        # Student logs in and replies
        self.client.login(username='student_clara', password='Password123!')
        reply_res = self.client.post(reverse('student:messages'), {
            'message': 'Sir, I have started working on Project 1 and pushed the code.',
            'link_url': 'https://github.com/clara/project1'
        }, follow=True)
        self.assertEqual(reply_res.status_code, 200)
        self.client.logout()

        # Verify reply exists in database
        reply_msg = StudentDirectMessage.objects.filter(
            student=self.student,
            trainer=self.trainer,
            sender_role='STUDENT'
        ).first()
        self.assertIsNotNone(reply_msg)
        self.assertIn("pushed the code", reply_msg.message)
        self.assertEqual(reply_msg.sender_role, 'STUDENT')

        # Trainer logs in and views the 1-on-1 thread
        self.client.login(username='trainer_mark', password='Password123!')
        thread_res = self.client.get(reverse('trainer:student_conversation', kwargs={'student_id': self.student.id}))
        self.assertEqual(thread_res.status_code, 200)
        self.assertContains(thread_res, "pushed the code")
        self.assertContains(thread_res, "https://github.com/clara/project1")
        self.client.logout()

    def test_salary_deductions_full_day_and_half_day_leaves(self):
        """
        Verify salary deduction formula:
        Per day = 60,000 / 30 = 2,000.
        28 days fully present.
        1 Full Day leave (1.0 day deduction) -> ₹2,000.
        1 Half Day leave (0.5 day deduction) -> ₹1,000.
        1 Half Day punch penalty (>20m delay, 0.5 day deduction) -> ₹1,000.
        Total deductions = 2.0 days = ₹4,000.
        Net salary = 60,000 - 4,000 = ₹56,000.
        """
        from attendance.models import LeaveRequest, TrainerAttendance

        self.trainer.monthly_salary = Decimal("60000.00")
        self.trainer.save()

        # 28 days Present in April 2026
        for day in range(1, 29):
            TrainerAttendance.objects.create(
                trainer=self.trainer,
                date=date(2026, 4, day),
                status='Present'
            )

        # 1 Half Day punch penalty on April 29
        TrainerAttendance.objects.create(
            trainer=self.trainer,
            date=date(2026, 4, 29),
            status='Half_Day',
            remark='Late Punch-In by 25 mins (>20m grace) -> Half Day marked'
        )

        # Approved Full Day Leave on April 30
        LeaveRequest.objects.create(
            trainer=self.trainer,
            leave_duration_type='FULL_DAY',
            from_date=date(2026, 4, 30),
            to_date=date(2026, 4, 30),
            reason="Personal travel",
            status="Approved"
        )

        # Approved Half Day Leave on April 15
        LeaveRequest.objects.create(
            trainer=self.trainer,
            leave_duration_type='HALF_DAY',
            half_day_period='FIRST_HALF',
            from_date=date(2026, 4, 15),
            to_date=date(2026, 4, 15),
            reason="Doctor consultation",
            status="Approved"
        )

        self.client.login(username='admin_ram', password='Password123!')
        res = self.client.post(reverse('branch_head:salary_calculator'), {
            'trainer': self.trainer.id,
            'month': 4,
            'year': 2026,
            'working_days': 30,
            'max_paid_leaves': 0,
            'calculate_preview': '1'
        })
        self.assertEqual(res.status_code, 200)
        calc = res.context['calc']
        self.assertIsNotNone(calc)
        self.assertEqual(calc['per_day_salary'], Decimal("2000.00"))
        self.assertEqual(calc['full_day_leaves'], 1.0)
        self.assertEqual(calc['half_day_leaves'], 1.0)
        self.assertEqual(calc['half_day_punches'], 1)
        # Total deduction days: 1.0 (full leave) + 0.5 (half leave) + 0.5 (half day punch) = 2.0
        self.assertEqual(calc['total_deduction_days'], Decimal("2.0"))
        self.assertEqual(calc['deduction'], Decimal("4000.00")) # 2.0 * 2000
        self.assertEqual(calc['net_salary'], Decimal("56000.00")) # 60000 - 4000
        self.client.logout()

    def test_calendar_builder_color_coding(self):
        """Verify attendance calendar utility labels days with exact color codes."""
        from attendance.calendar_utils import build_month_calendar_grid

        # April 2026: April 5 is Sunday, April 6 is Monday
        records_by_day = {
            1: {'status': 'Present', 'check_in': '09:05 AM'},
            2: {'status': 'Half_Day', 'remark': 'Late Punch-In'},
            3: {'status': 'Absent'},
            7: {'status': 'Leave'}
        }

        cal_grid = build_month_calendar_grid(2026, 4, records_by_day, today=date(2026, 4, 20))
        self.assertEqual(cal_grid['present_count'], 1)
        self.assertEqual(cal_grid['half_day_count'], 1)
        self.assertEqual(cal_grid['leave_count'], 1)

        # Check Sunday detection
        all_days = [d for week in cal_grid['weeks'] for d in week if not d.get('is_padding')]
        day5 = next(d for d in all_days if d['day'] == 5)
        self.assertTrue(day5['is_sunday'])
        self.assertEqual(day5['color_class'], 'cal-sunday')

        # Check Present day
        day1 = next(d for d in all_days if d['day'] == 1)
        self.assertEqual(day1['color_class'], 'cal-present')

    def test_public_home_trainers_listed_at_top_with_social_links(self):
        """Verify public home lists trainers at the top with full social profiles."""
        self.trainer.twitter = "https://x.com/mark_tech"
        self.trainer.instagram = "https://instagram.com/mark_teach"
        self.trainer.youtube = "https://youtube.com/@mark_academy"
        self.trainer.save()

        response = self.client.get(reverse('public_home'))
        self.assertEqual(response.status_code, 200)
        content = response.content.decode('utf-8')

        # Trainers section comes first
        trainer_idx = content.find("Our Expert Faculty & Trainers")
        student_idx = content.find("Our Enrolled Scholars")
        self.assertTrue(trainer_idx != -1 and student_idx != -1)
        self.assertLess(trainer_idx, student_idx, "Trainers section must be located above Students roster")

        # Social links present
        self.assertIn("https://x.com/mark_tech", content)
        self.assertIn("https://instagram.com/mark_teach", content)
        self.assertIn("https://youtube.com/@mark_academy", content)

    def test_branch_head_export_students_excel(self):
        """Verify Branch Head can export all students directory to a real Excel (.xlsx) file."""
        from io import BytesIO
        import openpyxl

        self.client.login(username='admin_ram', password='Password123!')
        response = self.client.get(reverse('branch_head:export_students_excel'))
        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response['Content-Type'],
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        self.assertIn('attachment; filename="DeviX_Students_Directory.xlsx"', response['Content-Disposition'])

        # Verify Excel workbook is valid and contains Clara Oswald
        wb = openpyxl.load_workbook(BytesIO(response.content))
        ws = wb["Students Directory"]
        self.assertEqual(ws["A1"].value, "DeviX Academy — Students Master Directory")
        # Check student in data rows
        found_student = False
        for r in range(5, ws.max_row + 1):
            if ws.cell(row=r, column=3).value == "Clara Oswald":
                found_student = True
                self.assertEqual(ws.cell(row=r, column=2).value, "DVX-2026-0901")
        self.assertTrue(found_student, "Exported Excel must contain Clara Oswald")

    def test_branch_head_student_attendance_report_html_and_excel(self):
        """Verify Branch Head student monthly attendance report HTML view and Excel export."""
        from io import BytesIO
        import openpyxl

        # Seed attendance for Clara
        Attendance.objects.create(student=self.student, date=date.today(), status='Present', marked_by=self.bh_user)

        self.client.login(username='admin_ram', password='Password123!')
        # HTML Matrix View
        html_resp = self.client.get(reverse('branch_head:student_attendance_report'))
        self.assertEqual(html_resp.status_code, 200)
        self.assertIn("Students Monthly Attendance Matrix", html_resp.content.decode('utf-8'))
        self.assertIn("Clara Oswald", html_resp.content.decode('utf-8'))

        # Excel Export
        excel_resp = self.client.get(reverse('branch_head:export_student_attendance_excel') + f'?month={date.today().month}&year={date.today().year}')
        self.assertEqual(excel_resp.status_code, 200)
        wb = openpyxl.load_workbook(BytesIO(excel_resp.content))
        self.assertIn("Attendance", wb.sheetnames[0])
        ws = wb.active
        self.assertIn("DeviX Academy — Student Monthly Attendance Matrix", ws["A1"].value)

    def test_trainer_export_my_students_excel(self):
        """Verify Trainer can export their assigned students list to Excel (.xlsx)."""
        from io import BytesIO
        import openpyxl

        self.client.login(username='trainer_mark', password='Password123!')
        response = self.client.get(reverse('trainer:export_my_students_excel'))
        self.assertEqual(response.status_code, 200)
        wb = openpyxl.load_workbook(BytesIO(response.content))
        ws = wb["Students Directory"]
        self.assertEqual(ws.cell(row=5, column=3).value, "Clara Oswald")

    def test_trainer_student_attendance_report_and_excel(self):
        """Verify Trainer student monthly attendance report HTML view and Excel export."""
        from io import BytesIO
        import openpyxl

        self.client.login(username='trainer_mark', password='Password123!')
        # HTML View
        resp = self.client.get(reverse('trainer:student_attendance_report'))
        self.assertEqual(resp.status_code, 200)
        self.assertIn("My Students Monthly Attendance Matrix", resp.content.decode('utf-8'))

        # Excel Export
        excel_resp = self.client.get(reverse('trainer:export_student_attendance_excel') + f'?month={date.today().month}&year={date.today().year}')
        self.assertEqual(excel_resp.status_code, 200)
        wb = openpyxl.load_workbook(BytesIO(excel_resp.content))
        self.assertIn("Attendance", wb.sheetnames[0])

